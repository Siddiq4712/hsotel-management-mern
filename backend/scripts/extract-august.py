"""Read the supplied August workbook without changing it."""
import openpyxl, json, hashlib
from pathlib import Path

p = Path(r'C:\Users\HARIRAAM\Documents\August 2026 mess bill  preparation sheet.xlsx')
w = openpyxl.load_workbook(p, data_only=True)
students = []
for name, college in [('MANDAYS', 'nec'), ('LAPC MANDAYS', 'lapc')]:
    for row in w[name].iter_rows(min_row=4):
        v = [c.value for c in row]
        if not isinstance(v[0], (int, float)) or not v[2] or not isinstance(v[3], (int, float)):
            continue
        nec = college == 'nec'
        students.append(dict(sheet=name, row=row[0].row, roll=str(v[1]).strip() if v[1] is not None else None,
            name=str(v[2]).strip(), college=college, days=v[3], rate=v[4], mess=v[5] or 0,
            additional=v[6] or 0, bed=(v[7] or 0) if nec else 0, newspaper=(v[9] or 0) if nec else 0,
            guest=(v[8] or 0) if nec else 0, fine=(v[10] or 0) if nec else 0,
            total=v[11] if nec else v[7], final=(v[12] if nec else v[8]) or 0))
s = w['DAILY RATE ']
data = dict(month=8, year=2026, source=str(p), sha256=hashlib.sha256(p.read_bytes()).hexdigest(), students=students,
    expenses=[dict(name=s.cell(r,2).value.strip(), amount=s.cell(r,3).value) for r in range(4,17)],
    deductions=[dict(name=n,amount=s.cell(r,3).value) for r,n in [(18,'Cash Token'),(19,'Sister Concern Bill'),(20,'Student Additional Credit Token'),(21,'Student Guest Income')]],
    sourceTotals=dict(gross=s['C17'].value,net=s['C22'].value,days=s['C24'].value,rate=s['C25'].value,final=w['MANDAYS']['M425'].value))
assert len(students) == 537
assert sum(x['days'] for x in students) == data['sourceTotals']['days']
assert sum(x['final'] for x in students) == data['sourceTotals']['final']
out = Path(__file__).resolve().parents[1] / 'import-work/august-2026'
out.mkdir(parents=True, exist_ok=True)
(out/'source.json').write_text(json.dumps(data, indent=2), encoding='utf8')
print(json.dumps(dict(rows=len(students), days=sum(x['days'] for x in students), bills=sum(x['final'] for x in students))))
